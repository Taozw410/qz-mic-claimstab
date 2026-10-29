from __future__ import annotations

import json
import re
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parent.parent
PROJECT_ROOT = ROOT.parent
MASTER_PATH = PROJECT_ROOT / "outputs" / "QZ-mic-data-master" / "QZ_mic_master_crosswalk.xlsx"
OAT_PATH = PROJECT_ROOT / "数据" / "燕麦生理指标.xlsx"
SAMPLE_SHEET_PATH = ROOT / "01_数据" / "sample_sheet.csv"
OUTPUT_DIR = ROOT / "03_结果" / "08_植物生理数据核查"

PRIMARY_SAMPLES = {
    "1_1",
    "1_2",
    "1_3",
    "20_1",
    "20_2",
    "20_3",
    "7_1",
    "7_2",
    "7_3",
    "22_1",
    "22_2",
    "22_3",
    "24_1",
    "24_2",
    "24_3",
}


def read_master() -> pd.DataFrame:
    workbook = load_workbook(MASTER_PATH, read_only=True, data_only=True)
    sheet = workbook["唯一主表"]
    rows = list(sheet.values)
    headers = list(rows[2])
    return pd.DataFrame(rows[3:], columns=headers).dropna(how="all")


def parse_replicate(label: object) -> int | None:
    if label is None:
        return None
    match = re.search(r"-([1-3])$", str(label))
    return int(match.group(1)) if match else None


def build_sample_sheet(master: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame]:
    required = [
        "唯一记录ID",
        "处理编号",
        "小区重复",
        "小区重复标签",
        "田间小区唯一ID",
        "组成代码",
        "密度等级",
        "原始测序组",
        "测序样名",
        "土壤记录处理号",
        "土壤原始样品号",
        "映射状态",
    ]
    missing = [column for column in required if column not in master.columns]
    if missing:
        raise ValueError(f"Unique master is missing columns: {missing}")

    sequenced = master.loc[master["测序样名"].notna()].copy()
    sequenced["Primary5"] = sequenced["测序样名"].isin(PRIMARY_SAMPLES)
    sequenced["ReplicateFromLabel"] = sequenced["小区重复"].map(parse_replicate)
    confirmed_random_id = sequenced["田间小区唯一ID"].notna()

    sample_sheet = pd.DataFrame(
        {
            "Sample": sequenced["测序样名"],
            "RawGroup": sequenced["原始测序组"],
            "Treatment": sequenced["处理编号"],
            "Composition": sequenced["组成代码"].fillna("Unassigned"),
            "Density": sequenced["密度等级"].fillna("Unknown"),
            "Replicate": sequenced["ReplicateFromLabel"],
            "PlotRecordID": sequenced["唯一记录ID"],
            "PlotRepeatLabel": sequenced["小区重复标签"],
            "FieldPlotID": sequenced["田间小区唯一ID"],
            "SoilRecordID": sequenced["土壤记录处理号"],
            "SoilSourceSampleID": sequenced["土壤记录处理号"],
            "SoilMappingLevel": np.where(
                sequenced["土壤记录处理号"].notna(), "treatment", "unavailable"
            ),
            "Primary5": sequenced["Primary5"],
            "SourceMappingStatus": sequenced["映射状态"],
            "MappingStatus": np.where(
                confirmed_random_id,
                "已确认：FieldPlotID为用户确认的正式田间随机排列号",
                sequenced["映射状态"],
            ),
        }
    ).sort_values("Sample", key=lambda values: values.map(lambda value: (
        int(str(value).split("_")[0]),
        int(str(value).split("_")[1]),
    )))
    sample_sheet["Replicate"] = sample_sheet["Replicate"].astype("Int64")

    primary = sample_sheet.loc[sample_sheet["Primary5"]].copy()
    if len(primary) != 15:
        raise ValueError(f"Expected 15 primary samples, found {len(primary)}")
    if primary["PlotRecordID"].nunique() != 15:
        raise ValueError("Primary plot records are not unique")
    counts = primary.groupby("Composition").size()
    if not counts.eq(3).all():
        raise ValueError(f"Expected three primary records per group: {counts.to_dict()}")

    all_plots = master.loc[master["田间小区唯一ID"].notna()].copy()
    if len(all_plots) != 69:
        raise ValueError(f"Expected 69 confirmed plot rows, found {len(all_plots)}")
    if all_plots["田间小区唯一ID"].nunique() != 69:
        raise ValueError("FieldPlotID is not unique across the 69 confirmed plot rows")

    audit = pd.DataFrame(
        [
            {"Check": "confirmed_plot_rows", "Value": len(all_plots), "Status": "PASS"},
            {
                "Check": "unique_FieldPlotID",
                "Value": all_plots["田间小区唯一ID"].nunique(),
                "Status": "PASS",
            },
            {"Check": "primary_sequenced_rows", "Value": len(primary), "Status": "PASS"},
            {
                "Check": "primary_unique_PlotRecordID",
                "Value": primary["PlotRecordID"].nunique(),
                "Status": "PASS",
            },
            {
                "Check": "explicit_block_column_present",
                "Value": int(any("区组" in str(column) for column in master.columns)),
                "Status": "FAIL",
            },
            {
                "Check": "source_mapping_status_says_random_id_pending",
                "Value": int(
                    primary["SourceMappingStatus"]
                    .fillna("")
                    .str.contains("真实随机排列ID待核对", regex=False)
                    .sum()
                ),
                "Status": "WARN",
            },
            {
                "Check": "derived_primary_mapping_status_confirmed",
                "Value": int(
                    sample_sheet.loc[sample_sheet["Primary5"], "MappingStatus"]
                    .str.startswith("已确认", na=False)
                    .sum()
                ),
                "Status": "PASS",
            },
        ]
    )
    return sample_sheet, audit


def treatment_design(raw_master: pd.DataFrame) -> pd.DataFrame:
    workbook = load_workbook(MASTER_PATH, read_only=True, data_only=True)
    sheet = workbook["处理设计"]
    rows = list(sheet.values)
    headers = list(rows[2])
    design = pd.DataFrame(rows[3:], columns=headers).dropna(how="all")
    return design.set_index("处理编号")


def normalize_oat_treatment(value: object) -> str:
    return str(value).strip()


def read_oat_metrics() -> tuple[pd.DataFrame, list[str]]:
    workbook = load_workbook(OAT_PATH, read_only=True, data_only=True)

    weight_sheet = workbook["重量"]
    weight_rows = list(weight_sheet.values)
    weight_data = []
    for row in weight_rows[1:13]:
        treatment = normalize_oat_treatment(row[0])
        wet = pd.to_numeric(pd.Series(row[3:13]), errors="coerce").dropna()
        weight_data.append(
            {
                "Treatment": treatment,
                "DryWeight5Plants": pd.to_numeric(row[1], errors="coerce"),
                "WetWeightNPlants": len(wet),
                "WetWeightMean": wet.mean(),
                "WetWeightSD": wet.std(ddof=1),
            }
        )

    height_sheet = workbook["高度"]
    height_rows = list(height_sheet.values)
    height_data = []
    for row in height_rows[1:13]:
        treatment = normalize_oat_treatment(row[0])
        heights = pd.to_numeric(pd.Series(row[1:11]), errors="coerce").dropna()
        height_data.append(
            {
                "Treatment": treatment,
                "HeightNPlants": len(heights),
                "HeightMean": heights.mean(),
                "HeightSD": heights.std(ddof=1),
            }
        )

    oat = pd.DataFrame(weight_data).merge(
        pd.DataFrame(height_data), on="Treatment", how="outer", validate="one_to_one"
    )
    excluded = [
        "重量!B2:B13: 5株总干重，每个处理只有一个处理级汇总值",
        "重量!D2:M13 与高度!B2:K13: 每处理10株，是植株子样本，不是小区重复",
        "高度!A32:G55: 工作表内单因素方差分析按10株作为重复，未用于论文推断",
        "Sheet1!A2:D11: 仅有p1-1至p1-10，无法对应本研究处理，排除",
    ]
    return oat, excluded


def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    master = read_master()
    sample_sheet, audit = build_sample_sheet(master)
    design = treatment_design(master)
    oat, excluded = read_oat_metrics()

    design_columns = [
        "组成代码",
        "物种组成",
        "密度等级",
        "燕麦/m²",
        "垂穗披碱草/m²",
        "草地早熟禾/m²",
        "中华羊茅/m²",
    ]
    oat = oat.merge(
        design[design_columns],
        left_on="Treatment",
        right_index=True,
        how="left",
        validate="one_to_one",
    )
    oat["AnalysisLevel"] = "treatment"
    oat["PlotLevelReplicates"] = 0
    oat["InferentialUse"] = "descriptive context only; plant rows are subsamples"
    oat.to_csv(OUTPUT_DIR / "oat_physiology_treatment_level.csv", index=False)

    link = sample_sheet.loc[
        sample_sheet["Primary5"],
        [
            "Sample",
            "Treatment",
            "Composition",
            "Density",
            "PlotRecordID",
            "PlotRepeatLabel",
            "FieldPlotID",
        ],
    ].copy()
    link["OatDataAvailable"] = link["Treatment"].isin(oat["Treatment"])
    link["OatMappingLevel"] = np.where(link["OatDataAvailable"], "treatment", "unavailable")
    link["CanPairWithPlotLevelOat"] = False
    link.to_csv(OUTPUT_DIR / "micro_plot_to_oat_link.csv", index=False)

    sample_sheet.to_csv(SAMPLE_SHEET_PATH, index=False)
    audit.to_csv(OUTPUT_DIR / "crosswalk_audit.csv", index=False)

    summary = {
        "sample_sheet_rows": len(sample_sheet),
        "primary_rows": int(sample_sheet["Primary5"].sum()),
        "oat_treatments": oat["Treatment"].tolist(),
        "microbial_primary_groups_with_oat_treatment_data": (
            link.loc[link["OatDataAvailable"], "Composition"].drop_duplicates().tolist()
        ),
        "microbial_primary_groups_without_oat_treatment_data": (
            link.loc[~link["OatDataAvailable"], "Composition"].drop_duplicates().tolist()
        ),
        "excluded_oat_ranges": excluded,
    }
    (OUTPUT_DIR / "audit_summary.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
