from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent.parent
MASTER = ROOT / "05_稿件" / "10_MASTER_draft_en.md"
OUTPUT = ROOT / "05_稿件" / "通用版.docx"

BODY_FONT = "Times New Roman"
HEADING_FONT = "Arial"
BODY_SIZE = 10.5
TABLE_HEADER_FILL = "1F4E79"
TABLE_ALT_FILL = "EAF1F8"
TABLE_BORDER = "D9D9D9"

HEADING_1 = {
    "Abstract",
    "Introduction",
    "Materials and methods",
    "Results",
    "Discussion",
    "Conclusion",
    "References",
    "Declarations",
}


def set_run_font(run, name: str, size: float | None = None, bold: bool | None = None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)


def add_hyperlink(paragraph, text: str, url: str, size: float = BODY_SIZE):
    part = paragraph.part
    relationship_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    new_run = OxmlElement("w:r")
    run_properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1F4E79")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), BODY_FONT)
    fonts.set(qn("w:hAnsi"), BODY_FONT)
    font_size = OxmlElement("w:sz")
    font_size.set(qn("w:val"), str(int(round(size * 2))))
    run_properties.extend([fonts, color, underline, font_size])
    new_run.append(run_properties)
    text_element = OxmlElement("w:t")
    text_element.text = text
    new_run.append(text_element)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def add_inline_markdown(paragraph, text: str, size: float = BODY_SIZE):
    token_pattern = re.compile(
        r"(`[^`]+`|\*\*.*?\*\*|\*.*?\*|https?://\S+)",
        flags=re.DOTALL,
    )
    position = 0
    for match in token_pattern.finditer(text):
        if match.start() > position:
            run = paragraph.add_run(text[position : match.start()])
            set_run_font(run, BODY_FONT, size)
        token = match.group(0)
        if token.startswith("`") and token.endswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, "Courier New", max(8.5, size - 0.8))
        elif token.startswith("**") and token.endswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, BODY_FONT, size, bold=True)
        elif token.startswith("*") and token.endswith("*"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, BODY_FONT, size)
            run.italic = True
        else:
            trailing = ""
            while token and token[-1] in ".,;)":
                trailing = token[-1] + trailing
                token = token[:-1]
            add_hyperlink(paragraph, token, token, size=size)
            if trailing:
                run = paragraph.add_run(trailing)
                set_run_font(run, BODY_FONT, size)
        position = match.end()
    if position < len(text):
        run = paragraph.add_run(text[position:])
        set_run_font(run, BODY_FONT, size)


def set_cell_margins(cell, top=70, start=90, bottom=70, end=90):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin_name, margin_value in (
        ("top", top),
        ("start", start),
        ("bottom", bottom),
        ("end", end),
    ):
        node = tc_mar.find(qn(f"w:{margin_name}"))
        if node is None:
            node = OxmlElement(f"w:{margin_name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(margin_value))
        node.set(qn("w:type"), "dxa")


def shade_cell(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = TABLE_BORDER, size: int = 4):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), str(size))
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    table_header = OxmlElement("w:tblHeader")
    table_header.set(qn("w:val"), "true")
    tr_pr.append(table_header)


def set_fixed_table_grid(table, widths_inches: list[float]):
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    table_width = OxmlElement("w:tblW")
    table_width.set(qn("w:w"), str(int(round(sum(widths_inches) * 1440))))
    table_width.set(qn("w:type"), "dxa")
    existing_width = tbl_pr.find(qn("w:tblW"))
    if existing_width is not None:
        tbl_pr.remove(existing_width)
    tbl_pr.append(table_width)
    grid_columns = table._tbl.tblGrid.findall(qn("w:gridCol"))
    for grid_column, width in zip(grid_columns, widths_inches):
        grid_column.set(qn("w:w"), str(int(round(width * 1440))))


def set_row_cant_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    set_run_font(run, HEADING_FONT, 8)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, text, end])


def configure_section(section, landscape: bool = False):
    section.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
    if landscape:
        section.page_width = Cm(29.7)
        section.page_height = Cm(21.0)
        section.left_margin = Cm(1.7)
        section.right_margin = Cm(1.7)
        section.top_margin = Cm(1.7)
        section.bottom_margin = Cm(1.7)
    else:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.left_margin = Cm(2.3)
        section.right_margin = Cm(2.3)
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
    section.header_distance = Cm(1.0)
    section.footer_distance = Cm(1.0)


def set_image_alt(inline_shape, title: str, description: str):
    doc_pr = inline_shape._inline.docPr
    doc_pr.set("title", title)
    doc_pr.set("descr", description)


def add_section(doc: Document, landscape: bool):
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    configure_section(section, landscape=landscape)
    section.footer.is_linked_to_previous = True
    return section


def parse_markdown_table(lines: list[str], start: int):
    rows = []
    index = start
    while index < len(lines) and lines[index].strip().startswith("|"):
        line = lines[index].strip()
        if not re.fullmatch(r"\|(?:\s*:?-+:?\s*\|)+", line):
            cells = [cell.strip() for cell in line.strip("|").split("|")]
            rows.append(cells)
        index += 1
    return rows, index


def infer_alignments(separator_line: str, columns: int):
    cells = [cell.strip() for cell in separator_line.strip("|").split("|")]
    alignments = []
    for cell in cells[:columns]:
        if cell.startswith(":") and cell.endswith(":"):
            alignments.append(WD_ALIGN_PARAGRAPH.CENTER)
        elif cell.endswith(":"):
            alignments.append(WD_ALIGN_PARAGRAPH.RIGHT)
        else:
            alignments.append(WD_ALIGN_PARAGRAPH.LEFT)
    return alignments + [WD_ALIGN_PARAGRAPH.LEFT] * (columns - len(alignments))


def add_table(doc: Document, rows: list[str], alignments, landscape: bool):
    columns = len(rows[0])
    table = doc.add_table(rows=len(rows), cols=columns)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    available_width = 9.9 if landscape else 6.5
    headers = [re.sub(r"[`*]", "", value).strip() for value in rows[0]]
    special_widths = None
    if columns == 13 and headers[0] == "Domain":
        special_widths = [
            0.55,
            0.80,
            0.60,
            0.70,
            0.75,
            0.45,
            0.55,
            0.70,
            0.75,
            0.45,
            0.55,
            0.70,
            0.65,
        ]
    elif columns == 11 and headers[0] == "Domain":
        special_widths = [
            0.65, 0.95, 0.75, 0.85, 0.85, 0.85, 0.85, 0.95, 0.90, 0.65, 0.65
        ]
    elif columns == 9 and headers[0] == "Domain":
        special_widths = [0.78, 1.08, 0.82, 1.48, 1.38, 1.48, 1.38, 1.52, 1.18]
    elif columns == 8 and headers[0] == "claim_id":
        special_widths = [1.18, 1.55, 0.90, 1.15, 1.10, 1.02, 0.92, 2.48]
    elif columns == 6 and headers[0] == "ID":
        special_widths = [0.55, 2.00, 1.15, 1.45, 2.65, 2.05]
    elif columns == 6 and headers[0] == "Domain":
        special_widths = [0.70, 2.05, 1.35, 1.50, 1.35, 2.65]
    elif columns == 7 and headers[0] == "Dataset":
        special_widths = [1.15, 1.65, 0.58, 0.92, 0.88, 0.98, 1.35]
    elif columns == 7 and headers[0] == "Dataset and claim":
        special_widths = [2.15, 0.95, 0.90, 0.80, 0.90, 0.75, 3.45]
    if special_widths:
        scale = available_width / sum(special_widths)
        widths = [value * scale for value in special_widths]
    else:
        weights = []
        for col in range(columns):
            lengths = [
                len(re.sub(r"[`*]", "", row[col]))
                for row in rows
                if col < len(row)
            ]
            weights.append(max(7, min(42, max(lengths) if lengths else 7)))
        total_weight = sum(weights)
        widths = [available_width * weight / total_weight for weight in weights]
    font_size = (
        6.6
        if columns >= 10
        else 7.3
        if columns >= 8
        else 7.8
        if columns >= 6
        else 8.2
    )
    set_fixed_table_grid(table, widths)

    for row_index, row_values in enumerate(rows):
        row = table.rows[row_index]
        set_row_cant_split(row)
        if row_index == 0:
            set_repeat_table_header(row)
        for col_index in range(columns):
            cell = row.cells[col_index]
            cell.width = Inches(widths[col_index])
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)
            set_cell_border(cell)
            if row_index == 0:
                shade_cell(cell, TABLE_HEADER_FILL)
            elif row_index % 2 == 0:
                shade_cell(cell, TABLE_ALT_FILL)
            paragraph = cell.paragraphs[0]
            paragraph.alignment = (
                alignments[col_index]
                if col_index < len(alignments)
                else WD_ALIGN_PARAGRAPH.LEFT
            )
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.0
            value = row_values[col_index] if col_index < len(row_values) else ""
            add_inline_markdown(paragraph, value, size=font_size)
            for run in paragraph.runs:
                if row_index == 0:
                    run.font.color.rgb = RGBColor(255, 255, 255)
                    run.bold = True
                set_run_font(run, BODY_FONT, font_size, bold=run.bold)
    return table


def add_para(doc: Document, text: str, style: str | None = None):
    paragraph = doc.add_paragraph(style=style)
    paragraph.paragraph_format.space_after = Pt(5)
    paragraph.paragraph_format.line_spacing = 1.12
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    add_inline_markdown(paragraph, text)
    return paragraph


def add_caption(doc: Document, text: str):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.keep_with_next = text.startswith("Table")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    add_inline_markdown(paragraph, text, size=9.2)
    return paragraph


def add_figure(doc: Document, path: Path, alt: str, landscape: bool):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.keep_with_next = True
    width = Inches(9.7 if landscape else 6.35)
    inline_shape = paragraph.add_run().add_picture(str(path), width=width)
    set_image_alt(inline_shape, alt, alt)


def configure_styles(doc: Document):
    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), BODY_FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), BODY_FONT)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    normal.font.size = Pt(BODY_SIZE)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.12

    for style_name, size, before, after in (
        ("Title", 21, 0, 12),
        ("Heading 1", 15, 12, 6),
        ("Heading 2", 12.5, 10, 4),
    ):
        style = doc.styles[style_name]
        style.font.name = HEADING_FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), HEADING_FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), HEADING_FONT)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), HEADING_FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        paragraph_properties = style._element.get_or_add_pPr()
        for border in paragraph_properties.findall(qn("w:pBdr")):
            paragraph_properties.remove(border)


def build_title_page(doc: Document, title: str):
    for _ in range(5):
        doc.add_paragraph()
    title_paragraph = doc.add_paragraph(style="Title")
    title_paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title_run = title_paragraph.add_run(title)
    set_run_font(title_run, HEADING_FONT, 21, bold=True)
    title_properties = title_paragraph._p.get_or_add_pPr()
    for border in title_properties.findall(qn("w:pBdr")):
        title_properties.remove(border)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_before = Pt(10)
    subtitle.paragraph_format.space_after = Pt(24)
    run = subtitle.add_run("General version with integrated figures and tables")
    set_run_font(run, HEADING_FONT, 13, bold=False)

    for text in (
        "Prepared 29 October 2026",
        (
            "Integrated contents: Abstract, main text, Figures 1–3, Figure S1, "
            "Tables 1–5, Tables S1–S4, references and declarations."
        ),
    ):
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(5)
        add_inline_markdown(paragraph, text, size=10)

    doc.add_page_break()


def main():
    lines = MASTER.read_text(encoding="utf-8").splitlines()
    title = next(
        line[2:].strip()
        for line in lines
        if line.startswith("# ")
    )

    doc = Document()
    configure_styles(doc)
    configure_section(doc.sections[0], landscape=False)
    add_page_number(doc.sections[0].footer.paragraphs[0])

    props = doc.core_properties
    props.title = title
    props.subject = "General version with integrated figures and tables"
    props.author = ""
    props.last_modified_by = ""
    props.comments = ""

    build_title_page(doc, title)

    current_landscape = False
    references_mode = False
    paragraph_buffer: list[str] = []
    i = 1

    def flush_paragraph():
        nonlocal paragraph_buffer
        if not paragraph_buffer:
            return
        text = " ".join(part.strip() for part in paragraph_buffer)
        text = re.sub(r"\s+", " ", text).strip()
        if text:
            paragraph = add_para(doc, text)
            if references_mode:
                paragraph.paragraph_format.left_indent = Cm(0.7)
                paragraph.paragraph_format.first_line_indent = Cm(-0.7)
                paragraph.paragraph_format.space_after = Pt(4)
                for run in paragraph.runs:
                    set_run_font(run, BODY_FONT, 9.0, bold=run.bold)
        paragraph_buffer = []

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if not stripped:
            flush_paragraph()
            i += 1
            continue

        if stripped.startswith("# "):
            flush_paragraph()
            heading = stripped[2:].strip()
            if heading in {"References", "Declarations"}:
                doc.add_page_break()
            references_mode = heading == "References"
            doc.add_heading(heading, level=1)
            i += 1
            continue

        if stripped.startswith("## "):
            flush_paragraph()
            heading = stripped[3:].strip()
            if heading.startswith("Table "):
                if heading.startswith("Table 1.") and not current_landscape:
                    add_section(doc, landscape=True)
                    current_landscape = True
                elif heading.startswith("Table S1.") and not current_landscape:
                    add_section(doc, landscape=True)
                    current_landscape = True
                elif heading.startswith("Table S3.") and current_landscape:
                    add_section(doc, landscape=False)
                    current_landscape = False
                elif any(
                    heading.startswith(prefix)
                    for prefix in (
                        "Table 2.",
                        "Table 3.",
                        "Table 4.",
                        "Table 5.",
                        "Table S2.",
                        "Table S4.",
                    )
                ):
                    doc.add_page_break()
                add_caption(doc, heading)
                i += 1
                continue
            if heading.startswith("Figure "):
                if heading == "Figure 2" and current_landscape:
                    add_section(doc, landscape=False)
                    current_landscape = False
                else:
                    doc.add_page_break()
                doc.add_heading(heading, level=1)
                i += 1
                continue
            if (
                heading in HEADING_1
                and heading != "Abstract"
                and heading != "Conclusion"
                and not references_mode
            ):
                doc.add_page_break()
            doc.add_heading(
                heading,
                level=1 if heading in HEADING_1 else 2,
            )
            i += 1
            continue

        if stripped.startswith("### "):
            flush_paragraph()
            heading = stripped[4:].strip()
            doc.add_heading(heading, level=2)
            i += 1
            continue

        image_match = re.fullmatch(r"!\[(.*?)\]\(<(.+?)>\)", stripped)
        if image_match:
            flush_paragraph()
            alt, path = image_match.groups()
            if re.match(r"^/[A-Za-z]:/", path):
                path = path[1:]
            add_figure(doc, Path(path), alt, current_landscape)
            i += 1
            continue

        if stripped.startswith("|"):
            flush_paragraph()
            rows, next_i = parse_markdown_table(lines, i)
            separator_line = lines[i + 1].strip() if i + 1 < len(lines) else ""
            alignments = infer_alignments(separator_line, len(rows[0]))
            add_table(doc, rows, alignments, current_landscape)
            i = next_i
            continue

        paragraph_buffer.append(stripped)
        i += 1

    flush_paragraph()

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
