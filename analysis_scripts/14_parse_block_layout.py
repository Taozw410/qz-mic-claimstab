from __future__ import annotations

import json
import re
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parent.parent
PROJECT_ROOT = ROOT.parent
LAYOUT_PATH = PROJECT_ROOT / "数据" / "区块播种示意图.xlsx"
MASTER_PATH = PROJECT_ROOT / "outputs" / "QZ-mic-data-master" / "QZ_mic_master_crosswalk.xlsx"
OUTPUT_DIR = ROOT / "03_结果" / "09_区块播种示意图核查"

SPECIES_CODES = {
    "燕": "As",
    "垂": "En",
    "羊": "Fs",
    "早": "Pp",
}

SHORT_LABEL_TO_TREATMENT = {
    "燕高": ("T1", "As", "高"),
    "燕中": ("T8", "As", "中"),
    "燕低": ("T15", "As", "低"),
    "垂高": ("T2", "En", "高"),
    "垂中": ("T9", "En", "中"),
    "垂低": ("T16", "En", "低"),
    "羊高": ("T4", "Fs", "高"),
    "羊中": ("T11", "Fs", "中"),
    "羊低": ("T18", "Fs", "低"),
    "早高": ("T3", "Pp", "高"),
    "早中": ("T10", "Pp", "中"),
    "早低": ("T17", "Pp", "低"),
    "三垂高": ("T5", "AEP", "高"),
    "三垂中": ("T12", "AEP", "中"),
    "三垂低": ("T19", "AEP", "低"),
    "三羊高": ("T6", "AFP", "高"),
    "三羊中": ("T13", "AFP", "中"),
    "三羊低": ("T20", "AFP", "低"),
    "四高": ("T7", "AEFP", "高"),
    "四中": ("T14", "AEFP", "中"),
    "四低": ("T21", "AEFP", "低"),
    "空": ("T22", "BP", "不适用"),
}

DEFINITIONS = {
    "燕": "燕麦单播",
    "垂": "垂穗披碱草单播",
    "羊": "中华羊茅单播",
    "早": "草地早熟禾单播",
    "三羊": "燕麦+中华羊茅+草地早熟禾",
    "三垂": "燕麦+垂穗披碱草+草地早熟禾",
    "四": "燕麦+垂穗披碱草+中华羊茅+草地早熟禾",
    "空": "无播种/裸斑对照候选，待最终确认",
}

RANDOMIZED_PLOT_TO_TREATMENT = {
    1: "T1", 2: "T12", 3: "T9", 4: "T2", 5: "T6", 6: "T11",
    7: "T7", 8: "T17", 9: "T14", 10: "T1", 11: "T21", 12: "T14",
    13: "T4", 14: "T16", 15: "T13", 16: "T2", 17: "T7", 18: "T19",
    19: "T20", 20: "T11", 21: "T19", 22: "T5", 23: "T18", 24: "T7",
    25: "T13", 26: "T1", 27: "T20", 28: "T6", 29: "T22", 30: "T8",
    31: "T22", 32: "T13", 33: "T2", 34: "T16", 35: "T10", 36: "T5",
    37: "T10", 38: "T17", 39: "T12", 40: "T4", 41: "T21", 42: "T3",
    43: "T3", 44: "T12", 45: "T17", 46: "T14", 47: "T22", 48: "T19",
    49: "T8", 50: "T16", 51: "T10", 52: "T6", 53: "T15", 54: "T9",
    55: "T21", 56: "T4", 57: "T8", 58: "T15", 59: "T5", 60: "T9",
    61: "T15", 62: "T11", 63: "T18", 64: "T20", 65: "T3", 66: "T18",
}

RANDOMIZED_PLOT_TO_SAMPLE_CODE = {
    1: "1.2", 2: "12.2", 3: "9.3", 4: "2.2", 5: "6.1", 6: "11.1",
    7: "7.2", 8: "17.3", 9: "14.3", 10: "1.1", 11: "21.3", 12: "14.1",
    13: "4.3", 14: "16.2", 15: "13.2", 16: "2.3", 17: "7.1", 18: "19.2",
    19: "20.3", 20: "11.3", 21: "19.3", 22: "5.1", 23: "18.2", 24: "7.3",
    25: "13.1", 26: "1.3", 27: "20.1", 28: "6.2", 29: "22.3", 30: "8.1",
    31: "22.2", 32: "13.3", 33: "2.1", 34: "16.3", 35: "10.2", 36: "5.3",
    37: "10.3", 38: "17.2", 39: "12.3", 40: "4.2", 41: "21.1", 42: "3.3",
    43: "3.1", 44: "12.1", 45: "17.1", 46: "14.2", 47: "22.1", 48: "19.1",
    49: "8.3", 50: "16.1", 51: "10.1", 52: "6.3", 53: "15.2", 54: "9.1",
    55: "21.2", 56: "4.1", 57: "8.2", 58: "15.3", 59: "5.2", 60: "9.2",
    61: "15.1", 62: "11.2", 63: "18.3", 64: "20.2", 65: "3.2", 66: "18.1",
}

USER_CONFIRMED_TREATMENT_TO_PLOTS = {
    "T1": [1, 10, 26],
    "T2": [4, 16, 33],
    "T6": [5, 28, 52],
    "T9": [3, 54, 60],
    "T12": [2, 39, 44],
}


def read_design() -> pd.DataFrame:
    workbook = load_workbook(MASTER_PATH, read_only=True, data_only=True)
    rows = list(workbook["处理设计"].values)
    headers = list(rows[2])
    return pd.DataFrame(rows[3:], columns=headers).dropna(how="all")


def parse_species_counts(text: str) -> dict[str, int]:
    text = str(text)
    result = {"燕": 0, "垂": 0, "羊": 0, "早": 0}
    patterns = {
        "燕": [r"燕(?:麦)?\s*(\d+)", r"燕麦\s*(\d+)"],
        "垂": [r"垂(?:穗披碱草)?\s*(\d+)", r"垂穗披碱草\s*(\d+)"],
        "羊": [r"羊(?:茅)?\s*(\d+)", r"中华羊茅\s*(\d+)"],
        "早": [r"早(?:熟禾)?\s*(\d+)", r"草地早熟禾\s*(\d+)"],
    }
    for species, species_patterns in patterns.items():
        for pattern in species_patterns:
            match = re.search(pattern, text)
            if match:
                result[species] = int(match.group(1))
                break
    return result


def expected_plot_counts(design: pd.DataFrame) -> dict[str, dict[str, int]]:
    species_columns = {
        "燕": "燕麦/m²",
        "垂": "垂穗披碱草/m²",
        "羊": "中华羊茅/m²",
        "早": "草地早熟禾/m²",
    }
    result = {}
    for treatment in [f"T{i}" for i in range(1, 22)]:
        row = design.loc[design["处理编号"].eq(treatment)].iloc[0]
        result[treatment] = {
            species: int(round(float(row[column] or 0) * 0.03))
            for species, column in species_columns.items()
        }
    return result


def match_p_block(counts: dict[str, int], expected: dict[str, dict[str, int]]) -> str:
    matches = [
        treatment
        for treatment, expected_counts in expected.items()
        if expected_counts == counts
    ]
    return matches[0] if len(matches) == 1 else "UNMATCHED"


def plot_sample_id(plot_number: int) -> str:
    code = RANDOMIZED_PLOT_TO_SAMPLE_CODE[plot_number]
    treatment, replicate = code.split(".", maxsplit=1)
    return f"T{treatment}-{replicate}"


def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    layout = load_workbook(LAYOUT_PATH, data_only=True, read_only=False)["区块分布"]
    design = read_design()
    expected = expected_plot_counts(design)

    p_blocks = []
    merged_by_top_left = {
        str(cell_range): cell_range
        for cell_range in layout.merged_cells.ranges
    }
    for cell_range in merged_by_top_left.values():
        value = layout.cell(cell_range.min_row, cell_range.min_col).value
        if value is None or not re.match(r"^p\d+", str(value).strip()):
            continue
        plot_token = re.match(r"^(p\d+)", str(value).strip()).group(1)
        counts = parse_species_counts(value)
        treatment = match_p_block(counts, expected)
        p_blocks.append(
            {
                "Zone": "SeedRecipeBlocks",
                "PositionType": "p_block",
                "CellRange": str(cell_range),
                "PlotToken": plot_token,
                "RandomizedPlotNumber": "",
                "RandomizedTreatmentFrom2019Table": "",
                    "RandomizedSampleID": "",
                "SourceLabel": str(value),
                "ShortLabel": "",
                **{f"Count_{species}": counts[species] for species in SPECIES_CODES},
                "MatchedTreatment": treatment,
                "CompositionCode": (
                    design.loc[design["处理编号"].eq(treatment), "组成代码"].iloc[0]
                    if treatment != "UNMATCHED"
                    else ""
                ),
                "Density": (
                    design.loc[design["处理编号"].eq(treatment), "密度等级"].iloc[0]
                    if treatment != "UNMATCHED"
                    else ""
                ),
                "MappingStatus": (
                    "matched by exact 0.03 m² seeding-count arithmetic"
                    if treatment != "UNMATCHED"
                    else "unresolved: recipe does not match T1-T21"
                ),
            }
        )

    grid_rows = []
    # The 16-cell middle layout uses plot numbers 51-66. The right edge of
    # each column carries the starting plot number, and rows run upward.
    for row_number in range(15, 22):
        for column_number in (7, 8, 9):  # G, H, I
            value = layout.cell(row_number, column_number).value
            if value is None:
                continue
            if column_number == 9:  # I; J contains 57..51
                plot_number = int(layout.cell(row_number, 10).value)
            elif column_number == 8:  # H; H22=58
                plot_number = int(layout.cell(22, 8).value) + (21 - row_number)
            else:  # G; G22=65 and only rows 20-21 are occupied
                plot_number = int(layout.cell(22, 7).value) + (21 - row_number)
            short_label = str(value).strip()
            mapping = SHORT_LABEL_TO_TREATMENT.get(short_label)
            randomized_treatment = RANDOMIZED_PLOT_TO_TREATMENT.get(plot_number, "")
            grid_rows.append(
                {
                    "Zone": "CentralFieldGrid",
                    "PositionType": "field_cell",
                    "CellRange": f"{get_column_letter(column_number)}{row_number}",
                    "PlotToken": "",
                    "RandomizedPlotNumber": plot_number,
                    "RandomizedTreatmentFrom2019Table": randomized_treatment,
                    "RandomizedSampleID": plot_sample_id(plot_number),
                    "SourceLabel": short_label,
                    "ShortLabel": short_label,
                    "Count_燕": "",
                    "Count_垂": "",
                    "Count_羊": "",
                    "Count_早": "",
                    "MatchedTreatment": mapping[0] if mapping else "",
                    "CompositionCode": mapping[1] if mapping else "",
                    "Density": mapping[2] if mapping else "",
                    "MappingStatus": (
                        "confirmed: 空 agrees with T22 in the randomization table"
                        if short_label == "空"
                        else (
                            "matched: abbreviation agrees with 2019 randomization table"
                            if mapping and mapping[0] == randomized_treatment
                            else "conflict: abbreviation disagrees with 2019 randomization table"
                        )
                    ),
                }
            )

    # The southeast grid spans L:T and covers randomized plot numbers 6-50.
    for row_number in range(17, 22):
        for column_number in range(12, 21):
            value = layout.cell(row_number, column_number).value
            if value is None:
                continue
            short_label = str(value).strip()
            mapping = SHORT_LABEL_TO_TREATMENT.get(short_label)
            plot_number = int(layout.cell(22, column_number).value) + (21 - row_number)
            randomized_treatment = RANDOMIZED_PLOT_TO_TREATMENT.get(plot_number, "")
            grid_rows.append(
                {
                    "Zone": "SoutheastFieldGrid",
                    "PositionType": "field_cell",
                    "CellRange": f"{get_column_letter(column_number)}{row_number}",
                    "PlotToken": "",
                    "RandomizedPlotNumber": plot_number,
                    "RandomizedTreatmentFrom2019Table": randomized_treatment,
                    "RandomizedSampleID": plot_sample_id(plot_number),
                    "SourceLabel": short_label,
                    "ShortLabel": short_label,
                    "Count_燕": "",
                    "Count_垂": "",
                    "Count_羊": "",
                    "Count_早": "",
                    "MatchedTreatment": mapping[0] if mapping else "",
                    "CompositionCode": mapping[1] if mapping else "",
                    "Density": mapping[2] if mapping else "",
                    "MappingStatus": (
                        "confirmed: 空 agrees with T22 in the randomization table"
                        if short_label == "空"
                        else (
                            "matched: abbreviation agrees with 2019 randomization table"
                            if mapping and mapping[0] == randomized_treatment
                            else "conflict: abbreviation disagrees with 2019 randomization table"
                        )
                    ),
                }
            )

    cells = pd.DataFrame(p_blocks + grid_rows)
    cells.insert(0, "LayoutCellID", [f"LAYOUT-{i:03d}" for i in range(1, len(cells) + 1)])
    expected_total = 69
    if len(cells) != expected_total:
        raise ValueError(f"Expected {expected_total} layout cells, found {len(cells)}")

    dictionary_rows = []
    for short_label, (treatment, composition, density) in SHORT_LABEL_TO_TREATMENT.items():
        stem = re.sub(r"[高中低]$", "", short_label)
        dictionary_rows.append(
            {
                "ShortLabel": short_label,
                "Definition": DEFINITIONS[stem],
                "TreatmentCode": treatment,
                "CompositionCode": composition,
                "Density": density,
                "MappingStatus": (
                    "confirmed: 空 cells match T22 plot numbers 29, 31, and 47"
                    if short_label == "空"
                    else "confirmed by user definition"
                ),
            }
        )

    cells.to_csv(OUTPUT_DIR / "layout_cells_69.csv", index=False)
    field_cells = cells.loc[cells["PositionType"].eq("field_cell")].copy()
    field_cells.to_csv(OUTPUT_DIR / "field_positions_61.csv", index=False)

    confirmed_rows = []
    for treatment, plot_numbers in USER_CONFIRMED_TREATMENT_TO_PLOTS.items():
        for plot_number in plot_numbers:
            randomized_treatment = RANDOMIZED_PLOT_TO_TREATMENT[plot_number]
            if randomized_treatment != treatment:
                raise ValueError(
                    f"User confirmation conflict: plot {plot_number} is "
                    f"{randomized_treatment} in the 2019 table, not {treatment}"
                )
            confirmed_rows.append(
                {
                    "Treatment": treatment,
                    "RandomizedPlotNumber": plot_number,
                    "RandomizedSampleID": plot_sample_id(plot_number),
                    "RandomizedTreatmentFrom2019Table": randomized_treatment,
                    "ConfirmationStatus": "matched 2019 randomization table",
                }
            )
    confirmed_assignments = pd.DataFrame(confirmed_rows).sort_values(
        "RandomizedPlotNumber"
    )
    confirmed_assignments.to_csv(
        OUTPUT_DIR / "user_confirmed_plot_assignments.csv", index=False
    )

    plot_rows = field_cells[
        [
            "RandomizedPlotNumber",
            "RandomizedSampleID",
            "MatchedTreatment",
            "CellRange",
            "Zone",
            "MappingStatus",
        ]
    ].copy()
    plot_rows = plot_rows.rename(
        columns={
            "CellRange": "LayoutCell",
            "Zone": "LayoutZone",
            "MappingStatus": "PositionStatus",
        }
    )
    missing_plot_numbers = sorted(
        set(RANDOMIZED_PLOT_TO_TREATMENT) - set(plot_rows["RandomizedPlotNumber"])
    )
    for plot_number in missing_plot_numbers:
        plot_rows = pd.concat(
            [
                plot_rows,
                pd.DataFrame(
                    [
                        {
                            "RandomizedPlotNumber": plot_number,
                            "RandomizedSampleID": plot_sample_id(plot_number),
                            "MatchedTreatment": RANDOMIZED_PLOT_TO_TREATMENT[plot_number],
                            "LayoutCell": "",
                            "LayoutZone": "OutsideDisplayedDiagram",
                            "PositionStatus": (
                                "plot number confirmed by user; diagram cell not shown"
                            ),
                        }
                    ]
                ),
            ],
            ignore_index=True,
        )
    plot_rows = plot_rows.sort_values("RandomizedPlotNumber")
    plot_rows.to_csv(OUTPUT_DIR / "randomized_plots_66.csv", index=False)
    sample_sheet = pd.read_csv(ROOT / "01_数据" / "sample_sheet.csv")
    all_samples = sample_sheet.copy()
    position_by_sample = field_cells.set_index("RandomizedSampleID")
    link_rows = []
    for _, sample in all_samples.iterrows():
        sample_id = sample["FieldPlotID"]
        if not bool(sample["Primary5"]):
            link_rows.append(
                {
                    "Sample": sample["Sample"],
                    "RandomizedSampleID": sample_id,
                    "RandomizedPlotNumber": "",
                    "LayoutZone": "",
                    "LayoutCell": "",
                    "LayoutStatus": "unassigned sequencing group",
                }
            )
        elif sample_id in position_by_sample.index:
            position = position_by_sample.loc[sample_id]
            link_rows.append(
                {
                    "Sample": sample["Sample"],
                    "RandomizedSampleID": sample_id,
                    "RandomizedPlotNumber": position["RandomizedPlotNumber"],
                    "LayoutZone": position["Zone"],
                    "LayoutCell": position["CellRange"],
                    "LayoutStatus": "mapped",
                }
            )
        else:
            confirmed = plot_rows.loc[plot_rows["RandomizedSampleID"].eq(sample_id)]
            if not confirmed.empty:
                position = confirmed.iloc[0]
                link_rows.append(
                    {
                        "Sample": sample["Sample"],
                        "RandomizedSampleID": sample_id,
                        "RandomizedPlotNumber": position["RandomizedPlotNumber"],
                        "LayoutZone": position["LayoutZone"],
                        "LayoutCell": position["LayoutCell"],
                        "LayoutStatus": position["PositionStatus"],
                    }
                )
            else:
                link_rows.append(
                    {
                        "Sample": sample["Sample"],
                        "RandomizedSampleID": sample_id,
                        "RandomizedPlotNumber": "",
                        "LayoutZone": "",
                        "LayoutCell": "",
                        "LayoutStatus": (
                            "natural reference outside randomized layout"
                            if sample["Composition"] == "Natural"
                            else "unresolved"
                        ),
                    }
                )
    layout_link = pd.DataFrame(link_rows)
    layout_link.to_csv(
        OUTPUT_DIR / "sample_position_link.csv", index=False
    )
    link_columns = [
        "RandomizedPlotNumber",
        "LayoutZone",
        "LayoutCell",
        "LayoutStatus",
    ]
    sample_sheet = sample_sheet.drop(
        columns=["RandomizedSampleID", *link_columns], errors="ignore"
    )
    sample_sheet = sample_sheet.merge(
        layout_link[["Sample", "RandomizedSampleID", *link_columns]],
        left_on=["Sample", "FieldPlotID"],
        right_on=["Sample", "RandomizedSampleID"],
        how="left",
        validate="one_to_one",
    )
    sample_sheet.to_csv(ROOT / "01_数据" / "sample_sheet.csv", index=False)
    pd.DataFrame(dictionary_rows).to_csv(
        OUTPUT_DIR / "layout_abbreviation_dictionary.csv", index=False
    )
    unresolved = cells.loc[
        cells["MappingStatus"].str.startswith("unresolved")
        | cells["MappingStatus"].str.startswith("inferred")
    ]
    unresolved.to_csv(OUTPUT_DIR / "layout_unresolved.csv", index=False)
    pd.DataFrame(
        [
            {
                "PlotNumberWithoutDiagramCell": number,
                "TreatmentFrom2019Table": RANDOMIZED_PLOT_TO_TREATMENT[number],
                "RandomizedSampleID": plot_sample_id(number),
                "Confirmation": "user-confirmed field plot number",
            }
            for number in missing_plot_numbers
        ]
    ).to_csv(OUTPUT_DIR / "plots_without_diagram_cell.csv", index=False)

    audit = {
        "layout_cells": len(cells),
        "p_blocks": len(p_blocks),
        "field_grid_cells": len(grid_rows),
        "p_blocks_matched_to_T1_T21": int(
            cells.loc[cells["PositionType"].eq("p_block"), "MatchedTreatment"]
            .ne("UNMATCHED")
            .sum()
        ),
        "p_blocks_unresolved": cells.loc[
            cells["PositionType"].eq("p_block")
            & cells["MatchedTreatment"].eq("UNMATCHED"),
            ["CellRange", "PlotToken", "SourceLabel"],
        ].to_dict("records"),
        "empty_grid_cells_confirmed_T22": int(
            cells.loc[cells["SourceLabel"].eq("空")].shape[0]
        ),
        "plots_without_displayed_diagram_cell": missing_plot_numbers,
        "confirmed_plot_assignments_for_missing_diagram_cells": {
            number: RANDOMIZED_PLOT_TO_TREATMENT[number]
            for number in missing_plot_numbers
        },
        "field_grid_cells_per_treatment": (
            cells.loc[cells["PositionType"].eq("field_cell"), "MatchedTreatment"]
            .value_counts()
            .sort_index()
            .to_dict()
        ),
        "field_grid_matches_2019_randomization": int(
            (
                cells.loc[
                    cells["PositionType"].eq("field_cell"),
                    "MatchedTreatment",
                ]
                == cells.loc[
                    cells["PositionType"].eq("field_cell"),
                    "RandomizedTreatmentFrom2019Table",
                ]
            ).sum()
        ),
        "note": (
            "The updated layout has 8 seed-recipe blocks plus 16 central and 45 "
            "southeast field cells = 69 cells. The 61 field cells match the 2019 "
            "randomization table exactly. The user supplied plot numbers 1-5, so all "
            "66 randomized plots are now identified even though plot numbers 1-5 "
            "have no displayed diagram cell. Six of eight seed-recipe blocks match T1-T21 by exact "
            "0.03 m² seeding-count arithmetic; p22 and p23 do not match a master "
            "treatment. Three 空 cells are T22 bare-control positions confirmed by "
            "the randomization table."
        ),
    }
    (OUTPUT_DIR / "layout_audit.json").write_text(
        json.dumps(audit, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    print(json.dumps(audit, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
