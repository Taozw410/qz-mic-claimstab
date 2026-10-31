from __future__ import annotations

import hashlib
import json
import re
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from pypdf import PdfReader


ROOT = Path(__file__).resolve().parent.parent
PACKAGE = ROOT / "投稿包_JMM_2026-10-30"
MANUSCRIPT = PACKAGE / "01_Manuscript" / "JMM_Manuscript_Anonymous.docx"
TITLE_PAGE = PACKAGE / "02_Title_Page" / "JMM_Title_Page.docx"
COVER_LETTER = PACKAGE / "03_Cover_Letter" / "JMM_Cover_Letter.docx"
SUPPLEMENT = (
    PACKAGE / "05_Supplementary" / "JMM_Supplementary_Material.docx"
)
DECLARATIONS = (
    PACKAGE
    / "06_Declarations_and_CRediT"
    / "JMM_Declarations_and_CRediT.docx"
)
REPORT = PACKAGE / "08_Submission_Audit" / "JMM_Package_Audit.md"
JSON_REPORT = (
    PACKAGE / "08_Submission_Audit" / "JMM_Package_Audit.json"
)
MANIFEST = PACKAGE / "08_Submission_Audit" / "Submission_Package_Manifest.md"


def word_count(text: str) -> int:
    return len(re.findall(r"\b[\w’'-]+\b", text))


def paragraph_after(doc: Document, heading: str, stop: str) -> str:
    paragraphs = [p.text.strip() for p in doc.paragraphs]
    start = paragraphs.index(heading)
    end = paragraphs.index(stop, start + 1)
    return " ".join(text for text in paragraphs[start + 1 : end] if text)


def xml_has_any(zf: zipfile.ZipFile, needle: str, files: list[str]):
    return any(
        needle in zf.read(name).decode("utf-8", errors="ignore")
        for name in files
        if name.endswith(".xml")
    )


def audit_docx(path: Path, manuscript: bool):
    doc = Document(path)
    result = {
        "file": path.name,
        "paragraphs": len(doc.paragraphs),
        "tables": len(doc.tables),
        "sections": len(doc.sections),
        "author": doc.core_properties.author,
        "last_modified_by": doc.core_properties.last_modified_by,
        "comments": doc.core_properties.comments,
    }
    with zipfile.ZipFile(path) as zf:
        names = zf.namelist()
        document_xml = zf.read("word/document.xml").decode("utf-8")
        result["has_page_field"] = any(
            " PAGE " in zf.read(name).decode("utf-8", errors="ignore")
            for name in names
            if re.fullmatch(r"word/footer\d+\.xml", name)
        )
        result["tracked_insertions"] = len(
            re.findall(r"<w:ins(?:\s|>)", document_xml)
        )
        result["tracked_deletions"] = len(
            re.findall(r"<w:del(?:\s|>)", document_xml)
        )
        result["comment_parts"] = [
            name for name in names if name.startswith("word/comments")
        ]
        result["drawing_count"] = document_xml.count("<w:drawing")
        result["line_number_rule_count"] = document_xml.count("<w:lnNumType")
        result["vertical_rule_count"] = sum(
            len(re.findall(r'<w:(?:left|right|insideV)[^>]*w:val="single"', block))
            for block in re.findall(r"<w:tcBorders>.*?</w:tcBorders>", document_xml)
        )
        result["cell_shading_count"] = document_xml.count("<w:shd")
        result["repeating_header_count"] = document_xml.count("<w:tblHeader")
        result["hidden_text_count"] = len(
            re.findall(r"<w:(?:vanish|webHidden)(?:\s|/|>)", document_xml)
        )
        result["embedded_object_count"] = len(
            re.findall(r"<w:object(?:\s|>)", document_xml)
        )
        result["embedded_parts"] = [
            name
            for name in names
            if name.startswith("word/embeddings/")
            or name.endswith("vbaProject.bin")
        ]

    if manuscript:
        abstract = paragraph_after(doc, "Abstract", "Keywords")
        keywords_text = paragraph_after(doc, "Keywords", "1. Introduction")
        all_text = "\n".join(p.text for p in doc.paragraphs)
        result["abstract_words"] = word_count(abstract)
        result["keywords"] = [
            value.strip()
            for value in keywords_text.split(";")
            if value.strip()
        ]
        result["numbered_sections"] = all(
            heading in all_text
            for heading in (
                "1. Introduction",
                "2. Materials and methods",
                "3. Results",
                "4. Discussion",
                "5. Conclusion",
                "6. Data and code availability",
                "7. Declarations",
            )
        )
        result["figure_captions_present"] = all(
            f"Figure {number}" in all_text for number in (1, 2, 3)
        )
        result["tables_cited"] = all(
            f"Table {number}" in all_text for number in range(1, 6)
        )
        result["release_urls_present"] = (
            "https://github.com/Taozw410/qz-mic-data/releases/tag/v1.0.0"
            in all_text
            and "https://github.com/Taozw410/qz-mic-claimstab/releases/tag/v0.3.0"
            in all_text
        )
    return result


def document_text(path: Path) -> str:
    return "\n".join(paragraph.text for paragraph in Document(path).paragraphs)


def clean_pdf_metadata(path: Path) -> dict:
    reader = PdfReader(path)
    metadata = dict(reader.metadata or {})
    return {
        "pages": len(reader.pages),
        "metadata": metadata,
        "contains_author_keys": any(
            key.lower() in {"/author", "/title", "/subject", "/keywords"}
            for key in metadata
        ),
    }


def main():
    manuscript = audit_docx(MANUSCRIPT, manuscript=True)
    title_page = audit_docx(TITLE_PAGE, manuscript=False)
    cover_letter = audit_docx(COVER_LETTER, manuscript=False)
    supplement = audit_docx(SUPPLEMENT, manuscript=False)
    declarations = audit_docx(DECLARATIONS, manuscript=False)
    title_text = document_text(TITLE_PAGE)
    cover_text = document_text(COVER_LETTER)
    declaration_text = document_text(DECLARATIONS)
    manuscript_text = document_text(MANUSCRIPT)

    figure_files = sorted((PACKAGE / "04_Figures").glob("*.pdf"))
    figure_headers = {
        path.name: path.read_bytes()[:5] for path in figure_files
    }
    figure_metadata = {
        path.name: clean_pdf_metadata(path) for path in figure_files
    }
    code_checksum_path = (
        PACKAGE / "07_Data_and_Code" / "SHA256SUMS.txt"
    )
    code_hashes_valid = True
    code_hash_rows = []
    for line in code_checksum_path.read_text(encoding="utf-8").splitlines():
        if not line.strip():
            continue
        expected, filename = line.split(None, 1)
        filename = filename.strip()
        target = package_path = (
            PACKAGE / "07_Data_and_Code" / filename
        )
        actual = hashlib.sha256(target.read_bytes()).hexdigest()
        matched = actual.lower() == expected.lower()
        code_hashes_valid = code_hashes_valid and matched
        code_hash_rows.append(
            {
                "file": filename,
                "expected": expected.lower(),
                "actual": actual,
                "matched": matched,
            }
        )

    temporary_markers = (
        ".tmp",
        ".temp",
        ".bak",
        "~$",
        ".DS_Store",
        "Thumbs.db",
    )
    package_files_initial = [
        path for path in PACKAGE.rglob("*") if path.is_file()
    ]
    secret_patterns = (
        re.compile(rb"ghp_[A-Za-z0-9]{20,}"),
        re.compile(rb"github_pat_[A-Za-z0-9_]{20,}"),
        re.compile(rb"BEGIN (?:RSA |OPENSSH |EC )?PRIVATE KEY"),
    )
    secret_hits = []
    for path in package_files_initial:
        data = path.read_bytes()
        if any(pattern.search(data) for pattern in secret_patterns):
            secret_hits.append(path.relative_to(PACKAGE).as_posix())

    a11y_reports = {}
    for path in sorted(
        (PACKAGE / "08_Submission_Audit").glob("a11y_*.json")
    ):
        a11y_reports[path.name] = json.loads(
            path.read_text(encoding="utf-8")
        )["counts"]

    technical_checks = {
        "abstract_under_250": manuscript["abstract_words"] <= 250,
        "keywords_1_to_7": 1 <= len(manuscript["keywords"]) <= 7,
        "numbered_sections": manuscript["numbered_sections"],
        "line_numbering_each_section": (
            manuscript["line_number_rule_count"] == manuscript["sections"]
            and supplement["line_number_rule_count"] == supplement["sections"]
        ),
        "page_field_present": (
            manuscript["has_page_field"] and supplement["has_page_field"]
        ),
        "no_tracked_changes": all(
            item["tracked_insertions"] == 0 and item["tracked_deletions"] == 0
            for item in (
                manuscript,
                title_page,
                cover_letter,
                supplement,
                declarations,
            )
        ),
        "no_comment_parts": all(
            not item["comment_parts"]
            for item in (
                manuscript,
                title_page,
                cover_letter,
                supplement,
                declarations,
            )
        ),
        "no_embedded_images_in_word_sources": (
            manuscript["drawing_count"] == 0
            and supplement["drawing_count"] == 0
        ),
        "docx_metadata_clean": all(
            not item["author"]
            and not item["last_modified_by"]
            and not item["comments"]
            for item in (
                manuscript,
                title_page,
                cover_letter,
                supplement,
                declarations,
            )
        ),
        "no_hidden_text": all(
            item["hidden_text_count"] == 0
            for item in (
                manuscript,
                title_page,
                cover_letter,
                supplement,
                declarations,
            )
        ),
        "no_embedded_objects_or_macros": all(
            item["embedded_object_count"] == 0
            and not item["embedded_parts"]
            for item in (
                manuscript,
                title_page,
                cover_letter,
                supplement,
                declarations,
            )
        ),
        "tables_horizontal_rules_only": (
            manuscript["vertical_rule_count"] == 0
            and supplement["vertical_rule_count"] == 0
        ),
        "tables_unshaded": (
            manuscript["cell_shading_count"] == 0
            and supplement["cell_shading_count"] == 0
        ),
        "table_headers_repeat": (
            manuscript["repeating_header_count"] >= 5
            and supplement["repeating_header_count"] >= 4
        ),
        "all_figures_are_pdf": (
            len(figure_files) == 4
            and all(header == b"%PDF-" for header in figure_headers.values())
        ),
        "figure_pdf_metadata_clean": all(
            not item["contains_author_keys"] for item in figure_metadata.values()
        ),
        "code_release_checksums_valid": code_hashes_valid,
        "no_temporary_or_editor_artifacts": not any(
            marker in path.name for marker in temporary_markers
            for path in package_files_initial
        ),
        "no_secret_patterns_in_package": not secret_hits,
        "accessibility_no_high_or_medium_issues": (
            sum(item.get("high", 0) for item in a11y_reports.values()) == 0
            and sum(item.get("medium", 0) for item in a11y_reports.values()) == 0
        ),
        "release_urls_in_manuscript": manuscript["release_urls_present"],
        "no_author_identity_in_anonymous_manuscript": not any(
            token in manuscript_text
            for token in (
                "taozw@gxust.edu.cn",
                "guangxixiejian@gxust.edu.cn",
                "489486943@qq.com",
                "0000-0003-0310-748X",
            )
        ),
        "manuscript_has_no_placeholders": not any(
            token in manuscript_text
            for token in (
                "AUTHOR TO",
                "TO BE INSERTED",
                "TELEPHONE REQUIRED",
            )
        ),
        "title_page_names_complete": all(
            name in title_text for name in ("ZiWei Tao", "Jian Xie", "Fan Yang")
        ),
        "title_page_affiliations_complete": (
            "2 Wenchang Road" in title_text
            and "508 No. 2 Street" in title_text
            and "545006" in title_text
            and "310018" in title_text
        ),
        "title_page_email_present": "taozw@gxust.edu.cn" in title_text,
        "title_page_orcid_present": "0000-0003-0310-748X" in title_text,
        "cover_letter_author_complete": (
            "ZiWei Tao" in cover_text
            and "taozw@gxust.edu.cn" in cover_text
            and "[CORRESPONDING AUTHOR" not in cover_text
        ),
        "declarations_complete": all(
            text in declaration_text
            for text in (
                "ZiWei Tao: Writing - original draft.",
                "Jian Xie: Data curation.",
                "Fan Yang: Conceptualization.",
                "no known competing financial interests",
                "did not receive any specific grant",
                "Generative AI tools were used to assist with language editing",
                "No new human participants or animals were involved.",
                "Not applicable.",
                "https://github.com/Taozw410/qz-mic-data/releases/tag/v1.0.0",
                "https://github.com/Taozw410/qz-mic-claimstab/releases/tag/v0.3.0",
            )
        ),
        "no_placeholders_in_title_cover_declarations": not any(
            token in text
            for text in (title_text, cover_text, declaration_text)
            for token in (
                "GIVEN NAME",
                "ADD OR REMOVE",
                "LIST FUNDER",
                "COMPLETE THE ELSEVIER",
                "[AUTHOR INITIALS]",
                "[TOOL/SERVICE]",
            )
        ),
    }

    author_action_checks = {
        "corresponding_author_telephone_supplied": (
            "TELEPHONE REQUIRED" not in title_text
            and re.search(r"Telephone:\s*\+?[0-9][0-9 ()-]{6,}", title_text)
            is not None
        )
    }

    self_generated = {REPORT, JSON_REPORT, MANIFEST}
    package_files = sorted(
        path
        for path in PACKAGE.rglob("*")
        if path.is_file() and path not in self_generated
    )
    inventory = [
        {
            "path": path.relative_to(PACKAGE).as_posix(),
            "bytes": path.stat().st_size,
            "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
        }
        for path in package_files
    ]

    payload = {
        "package": PACKAGE.name,
        "manuscript": manuscript,
        "title_page": title_page,
        "cover_letter": cover_letter,
        "supplement": supplement,
        "declarations": declarations,
        "figures": {
            name: {
                "pdf_header": header.decode("ascii", errors="replace"),
                "bytes": (PACKAGE / "04_Figures" / name).stat().st_size,
            }
            for name, header in figure_headers.items()
        },
        "code_checksums": code_hash_rows,
        "secret_hits": secret_hits,
        "accessibility_counts": a11y_reports,
        "technical_checks": technical_checks,
        "author_action_checks": author_action_checks,
        "overall_status": (
            "PASS"
            if all(technical_checks.values()) and all(author_action_checks.values())
            else "CONDITIONAL PASS"
        ),
        "files": inventory,
    }
    JSON_REPORT.write_text(
        json.dumps(payload, indent=2, ensure_ascii=False) + "\n",
        encoding="utf-8",
    )

    lines = [
        "# JMM submission package audit",
        "",
        f"- Package: `{PACKAGE.name}`",
        f"- Manuscript: {manuscript['paragraphs']} paragraphs, "
        f"{manuscript['tables']} tables, {manuscript['sections']} section",
        f"- Abstract: {manuscript['abstract_words']} words",
        f"- Keywords: {len(manuscript['keywords'])}",
        f"- Supplementary material: {supplement['tables']} tables, "
        f"{supplement['sections']} section",
        f"- Figure PDFs: {len(figure_files)}",
        "",
        "## Automated technical and hygiene checks",
        "",
        "| Check | Status |",
        "| --- | --- |",
    ]
    for name, passed in technical_checks.items():
        lines.append(f"| `{name}` | {'PASS' if passed else 'FAIL'} |")

    lines.extend(
        [
            "",
            "## Pre-submission author actions",
            "",
            "| Check | Status |",
            "| --- | --- |",
        ]
    )
    for name, passed in author_action_checks.items():
        lines.append(f"| `{name}` | {'PASS' if passed else 'ACTION REQUIRED'} |")

    lines.extend(
        [
            "",
            "- Enter the prepared competing-interest, funding, CRediT, and",
            "  generative-AI statements in Elsevier's declarations tool.",
            "- Confirm the final author order and all author approvals.",
            "- Select the subscription route if no APC is funded.",
            "- Optionally provide suggested and opposed reviewers.",
            "- If the institution requires a DOI in addition to the versioned",
            "  GitHub releases, archive them in Zenodo before submission.",
            "",
            "## Machine-readable report",
            "",
            "`JMM_Package_Audit.json` contains the complete file inventory and hashes.",
            "",
        ]
    )
    REPORT.write_text("\n".join(lines), encoding="utf-8")

    manifest_lines = [
        "# JMM submission package manifest",
        "",
        "Generated after the final package audit.",
        "",
        "| File | Size (bytes) | SHA256 |",
        "| --- | ---: | --- |",
    ]
    for path in sorted(
        path
        for path in PACKAGE.rglob("*")
        if path.is_file() and path != MANIFEST
    ):
        digest = hashlib.sha256(path.read_bytes()).hexdigest()
        manifest_lines.append(
            f"| `{path.relative_to(PACKAGE).as_posix()}` | "
            f"{path.stat().st_size:,} | `{digest}` |"
        )
    MANIFEST.write_text(
        "\n".join(manifest_lines) + "\n",
        encoding="utf-8",
    )

    technical_failed = [
        name for name, passed in technical_checks.items() if not passed
    ]
    author_failed = [
        name for name, passed in author_action_checks.items() if not passed
    ]
    if technical_failed:
        raise SystemExit(
            "Technical or hygiene audit failed: " + ", ".join(technical_failed)
        )
    if author_failed:
        print(
            "CONDITIONAL PASS: author actions remain: "
            + ", ".join(author_failed)
        )
    else:
        print("PASS")
    print(REPORT)


if __name__ == "__main__":
    main()
